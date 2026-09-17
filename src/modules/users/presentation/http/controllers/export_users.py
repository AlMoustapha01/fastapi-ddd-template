"""Export users HTTP controller."""

import contextlib
import csv
import io
from datetime import UTC, datetime
from enum import StrEnum
from inspect import getdoc
from typing import Annotated, Any

from dishka import FromDishka
from dishka.integrations.fastapi import inject
from fastapi import APIRouter, Depends, Query, status
from fastapi.responses import StreamingResponse
from fastapi_error_map import ErrorAwareRouter, rule

from core.lib.http.translators import (
    DomainErrorTranslator,
    ServiceUnavailableTranslator,
)
from modules.auth.application.dependencies import require_roles
from modules.auth.domain import AuthUser
from modules.users.application.queries import (
    ExportUserData,
    ExportUsersQuery,
    ExportUsersRequest,
)
from modules.users.domain.exceptions import UserValidationError

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False


class ExportFormat(StrEnum):
    """Export format options."""

    CSV = "csv"
    XLSX = "xlsx"


def create_export_users_router() -> APIRouter:
    """Create router for export users endpoint."""
    router = ErrorAwareRouter()

    @router.get(
        "/export",
        description=getdoc(ExportUsersQuery),
        error_map={
            UserValidationError: rule(
                status=status.HTTP_400_BAD_REQUEST,
                translator=DomainErrorTranslator(),
            ),
            Exception: rule(
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
                translator=ServiceUnavailableTranslator(),
            ),
        },
        status_code=status.HTTP_200_OK,
    )
    @inject
    async def export_users(
        query: FromDishka[ExportUsersQuery],
        _: Annotated[AuthUser, Depends(require_roles(["perm_view_user"]))],
        export_format: Annotated[
            ExportFormat,
            Query(description="Export format (csv or xlsx)", alias="format"),
        ] = ExportFormat.CSV,
        group_id: Annotated[
            str | None,
            Query(description="Filter by Keycloak group ID"),
        ] = None,
        is_active: Annotated[
            bool | None,
            Query(description="Filter by active status"),
        ] = None,
        search: Annotated[
            str | None,
            Query(description="Search by name, email, or username (case-insensitive)"),
        ] = None,
    ) -> StreamingResponse:
        """
        Export all users to CSV or XLSX format.

        Args:
            query: Injected ExportUsersQuery
            format: Export format (csv or xlsx)
            group_id: Filter by Keycloak group ID
            is_active: Filter by active status
            search: Search by name, email, or username

        Returns:
            StreamingResponse with file download

        Raises:
            UserValidationError: When validation fails
        """
        request_data = ExportUsersRequest(
            group_id=group_id,
            is_active=is_active,
            search=search,
        )
        users_data = await query.execute(request_data)

        # Generate filename with timestamp
        timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")

        if export_format == ExportFormat.CSV:
            return _export_to_csv(users_data, timestamp)
        if export_format == ExportFormat.XLSX:
            if not XLSX_AVAILABLE:
                raise ValueError("XLSX export requires openpyxl package")
            return _export_to_xlsx(users_data, timestamp)
        raise ValueError(f"Unsupported export format: {export_format}")

    return router


def _export_to_csv(
    users_data: list[ExportUserData], timestamp: str
) -> StreamingResponse:
    """Export users data to CSV format."""
    output = io.StringIO()

    if not users_data:
        writer: Any = csv.writer(output)
        writer.writerow([
            "ID",
            "Name",
            "Email",
            "Username",
            "Keycloak ID",
            "Status",
            "Created At",
            "Updated At",
            "Rôles",
            "Last Login",
            "Connected",
        ])
    else:
        fieldnames = [
            "id",
            "name",
            "email",
            "username",
            "keycloak_id",
            "is_active",
            "created_at",
            "updated_at",
            "groups",
            "last_login",
            "connected",
        ]
        writer = csv.DictWriter(output, fieldnames=fieldnames)

        # Write header with readable names
        writer.writerow({
            "id": "ID",
            "name": "Name",
            "email": "Email",
            "username": "Username",
            "keycloak_id": "Keycloak ID",
            "is_active": "Status",
            "created_at": "Created At",
            "updated_at": "Updated At",
            "groups": "Rôles",
            "last_login": "Last Login",
            "connected": "Connected",
        })

        # Write data rows
        writer.writerows(users_data)

    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename=users_export_{timestamp}.csv"
        },
    )


def _export_to_xlsx(
    users_data: list[ExportUserData], timestamp: str
) -> StreamingResponse:
    """Export users data to XLSX format."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Users"

    # Define headers
    headers = [
        "ID",
        "Name",
        "Email",
        "Username",
        "Keycloak ID",
        "Status",
        "Created At",
        "Updated At",
        "Rôles",
        "Last Login",
        "Connected",
    ]

    # Write headers with bold font
    for col_num, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num, value=header)
        cell.font = Font(bold=True)

    # Write data rows
    fieldnames = [
        "id",
        "name",
        "email",
        "username",
        "keycloak_id",
        "is_active",
        "created_at",
        "updated_at",
        "groups",
        "last_login",
        "connected",
    ]

    for row_num, user in enumerate(users_data, 2):
        for col_num, field in enumerate(fieldnames, 1):
            sheet.cell(row=row_num, column=col_num, value=user.get(field, ""))

    # Auto-adjust column widths
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            with contextlib.suppress(Exception):
                max_length = max(max_length, len(str(cell.value)))
        adjusted_width = min(max_length + 2, 50)
        sheet.column_dimensions[column_letter].width = adjusted_width

    # Save to bytes buffer
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=users_export_{timestamp}.xlsx"
        },
    )
